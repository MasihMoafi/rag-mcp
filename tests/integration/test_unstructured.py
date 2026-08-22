from pathlib import Path
import sys

from docx import Document
from openpyxl import Workbook
from PIL import Image, ImageDraw
from pptx import Presentation


sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from rag.core import RAGPipelineV2
from rag.fetch import find_all_indexable_files


def _pipeline() -> RAGPipelineV2:
    return RAGPipelineV2.__new__(RAGPipelineV2)


def test_office_and_csv_extraction(tmp_path: Path) -> None:
    docx_path = tmp_path / "sample.docx"
    document = Document()
    document.add_paragraph("DOCX multimodal acceptance 7319")
    document.save(docx_path)

    pptx_path = tmp_path / "sample.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    textbox = slide.shapes.add_textbox(0, 0, 8000000, 1000000)
    textbox.text = "PPTX multimodal acceptance 7319"
    presentation.save(pptx_path)

    xlsx_path = tmp_path / "sample.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["format", "value"])
    sheet.append(["XLSX multimodal acceptance", 7319])
    workbook.save(xlsx_path)

    csv_path = tmp_path / "sample.csv"
    csv_path.write_text(
        "format,value\nCSV multimodal acceptance,7319\n",
        encoding="utf-8",
    )

    pipeline = _pipeline()
    assert "DOCX multimodal acceptance 7319" in pipeline._extract_content(str(docx_path))
    assert "PPTX multimodal acceptance 7319" in pipeline._extract_content(str(pptx_path))
    assert "XLSX multimodal acceptance" in pipeline._extract_content(str(xlsx_path))
    assert "7319" in pipeline._extract_content(str(csv_path))


def test_image_ocr_extraction(tmp_path: Path) -> None:
    image_path = tmp_path / "sample.png"
    image = Image.new("RGB", (900, 180), "white")
    ImageDraw.Draw(image).text(
        (30, 60),
        "RAG multimodal test 7319",
        fill="black",
    )
    image.save(image_path)

    extracted = _pipeline()._extract_content(str(image_path))

    assert "7319" in extracted


def test_multimodal_extensions_are_discovered(tmp_path: Path) -> None:
    extensions = (".png", ".docx", ".pptx", ".xlsx", ".csv")
    for extension in extensions:
        (tmp_path / f"sample{extension}").touch()

    discovered = {
        Path(path).suffix
        for path in find_all_indexable_files(str(tmp_path))
    }

    assert set(extensions) <= discovered
